"""Pull pension fund data from the Central Bank of Iceland (Seðlabanki) into the raw zone.

Two kinds of Excel workbooks, both of which carry a raw long-format data sheet behind
their pivot-style report tabs:
  * the quarterly investment breakdown per fund and division ("Sundurliðun fjárfestinga",
    sheet Tegundaflokkun) → data/raw/sedlabanki/pension_investments.parquet
  * the annual summaries of the funds' financial statements (sheet Gögn), one workbook per
    year → data/raw/sedlabanki/pension_annual_<year>.parquet

A workbook whose bytes haven't changed since the last run is not re-parsed.
"""

from __future__ import annotations

import hashlib
import io
from datetime import datetime

from openpyxl import load_workbook

from common import RAW, PoliteSession, State, load_sources, run_safely, write_parquet

TARGET = RAW / "sedlabanki"

# Sheet Tegundaflokkun has no usable header row (row 1 is a pivot filter), so name columns by position.
# Columns: quarter-end (Excel date), the same as text, FundType, FundName, FundDivisionName,
# InvestmentTypeName, CurrencyCode, Bókfært virði (book value in ISK).
INVESTMENT_COLUMNS = ["quarter_end", "date_text", "fund_type", "fund_name", "division", "investment_type", "currency", "amount"]


def download(session: PoliteSession, base_url: str, itemid: str) -> bytes:
    resp = session.get(base_url, params={"itemid": itemid})
    resp.raise_for_status()
    if not resp.content.startswith(b"PK"):
        raise ValueError(f"item {itemid} is not an Excel file ({resp.headers.get('Content-Type')})")
    return resp.content


def read_sheet(content: bytes, sheet: str):
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not found (has {wb.sheetnames})")
    yield from wb[sheet].iter_rows(values_only=True)


def parse_investments(content: bytes, sheet: str) -> list[dict]:
    rows = []
    for values in read_sheet(content, sheet):
        # Data rows start with the quarter-end date; skip the filter and header rows above.
        if not values or not isinstance(values[0], datetime) or values[7] is None:
            continue
        row = dict(zip(INVESTMENT_COLUMNS, values[:8]))
        row["quarter_end"] = row["quarter_end"].date().isoformat()
        del row["date_text"]
        row["amount"] = float(row["amount"])
        rows.append(row)
    return rows


def parse_annual(content: bytes, sheet: str, year: int) -> list[dict]:
    it = read_sheet(content, sheet)
    header = None
    rows = []
    for values in it:
        if header is None:
            if values and "FundID" in values:
                header = [str(v) if v is not None else f"col{i}" for i, v in enumerate(values)]
            continue
        if not values or all(v is None for v in values):
            continue
        row = dict(zip(header, values))
        rows.append(
            {
                "year": year,
                "fund_id": row.get("FundID"),
                "fund_short_name": row.get("FundShortName"),
                "fund_name": row.get("FundName"),
                "division": row.get("FundDivisionName"),
                "fund_type": row.get("FundType"),
                "item_key": row.get("Lykill"),
                "section": row.get("Kafli"),
                "item_name": row.get("Name"),
                # 2025+ puts every figure in BokfaerdStada; earlier years put head counts in
                # Fjoldi and percentages in HlutfallsStada.
                "value": _first_num(row, "BokfaerdStada", "HlutfallsStada", "Fjoldi"),
                "net_assets": _num(row.get("HreinEign")),
                "pensions_paid": _num(row.get("Lífeyrir")),
            }
        )
    if header is None:
        raise ValueError("no header row with FundID found")
    return rows


def _first_num(row: dict, *columns: str):
    for column in columns:
        value = _num(row.get(column))
        if value is not None:
            return value
    return None


def _num(v):
    try:
        return float(v) if v is not None and v != "" else None
    except (TypeError, ValueError):
        return None


def main() -> bool:
    config = load_sources()["sedlabanki"]
    base_url = config["base_url"]
    state = State(TARGET / "_state.json")
    session = PoliteSession(min_interval=1.0)

    def fetch(name: str, itemid: str, parse) -> None:
        target = TARGET / f"{name}.parquet"
        content = download(session, base_url, itemid)
        digest = hashlib.sha256(content).hexdigest()
        if target.exists() and state.get(name) == digest:
            print(f"sedlabanki.{name}: unchanged, skipping")
            return
        rows = parse(content)
        write_parquet(rows, target, source_updated=digest[:12])
        state.set(name, digest)
        print(f"sedlabanki.{name}: {len(rows)} rows")

    inv = config["investments"]
    ok = run_safely(
        "sedlabanki.pension_investments",
        TARGET / "pension_investments.parquet",
        lambda: fetch("pension_investments", inv["itemid"], lambda c: parse_investments(c, inv["sheet"])),
    )
    annual = config["annual"]
    for year, itemid in sorted(annual["years"].items()):
        name = f"pension_annual_{year}"
        ok &= run_safely(
            f"sedlabanki.{name}",
            TARGET / f"{name}.parquet",
            lambda year=year, itemid=itemid, name=name: fetch(name, itemid, lambda c: parse_annual(c, annual["sheet"], int(year))),
        )
    return ok


if __name__ == "__main__":
    raise SystemExit(0 if main() else 1)
